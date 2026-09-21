from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)
NAVY = "16324F"
TEAL = "0F6B78"
LIGHT = "EAF2F5"
BORDER = "D9E2E8"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def font_run(run, name="Microsoft YaHei", size=10.5, bold=False, color="1F2937"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text: str, *, bold=False, size=10.5, color="1F2937"):
    run = paragraph.add_run(text)
    font_run(run, size=size, bold=bold, color=color)
    return run


def configure_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    doc.core_properties.title = title
    doc.core_properties.subject = "企业云产品智能支持与工单协同 Agent"


def add_title(doc: Document, title: str, subtitle: str, *, title_size: float = 24) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, title, bold=True, size=title_size, color="000000")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(12)
    add_text(sub, subtitle, size=10.5, color="4B5563")


def add_heading(doc: Document, text: str, level=1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(5)
    add_text(p, text, bold=True, size=16 if level == 1 else 12.5, color="000000")
    return p


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        run.text = ""
    add_text(p, text, size=10.2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Inches(widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, bold=True, size=9.5, color="FFFFFF")
    row_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    row_properties.append(repeat_header)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Inches(widths[idx])
            if row_index % 2:
                set_cell_shading(cell, "F5F9FB")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, str(value), size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def build_solution() -> Path:
    doc = Document()
    configure_document(doc, "企业云产品智能支持与工单协同 Agent 解决方案")
    add_title(doc, "企业云产品智能支持与工单协同 Agent", "两页解决方案  |  需求、架构、流程与验收")
    intro = doc.add_paragraph()
    add_text(
        intro,
        "本方案面向企业云 AI 产品咨询和 API 故障报修。系统先补齐故障信息，再检索文档和调用确定性工具；证据不足或风险较高时创建人工工单。",
        size=10.8,
    )
    add_heading(doc, "业务问题与方案目标")
    add_table(
        doc,
        ["环节", "传统处理痛点", "系统处理方式"],
        [
            ["问题受理", "描述不完整，反复沟通", "识别意图并主动收集错误码、时间和影响范围"],
            ["知识查询", "人工翻文档，答案缺少依据", "检索知识库并展示引用片段"],
            ["故障诊断", "模型可能猜测账户或服务状态", "Python 工具返回模拟额度、错误码和服务状态"],
            ["人工协同", "转交信息不完整，过程难追溯", "生成结构化工单并保留状态时间线"],
        ],
        [1.05, 2.7, 3.25],
    )
    add_heading(doc, "总体架构")
    image_path = ROOT / "docs" / "system_architecture.png"
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.2))

    second_page_heading = add_heading(doc, "端到端处理流程")
    second_page_heading.paragraph_format.page_break_before = True
    steps = [
        "输入脱敏：遮盖 API Key、SecretID、手机号和邮箱。",
        "分诊与追问：识别六类意图，信息不足时停止诊断并继续追问。",
        "检索与工具：知识库提供依据，Python 工具查询演示状态和额度。",
        "生成与升级：模型组织解决步骤；无依据、高风险或人工请求触发工单。",
        "处理与回流：工程师分派、解决、关闭，已验证方案进入候选知识区。",
    ]
    for index, step in enumerate(steps, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        add_text(p, f"{index}. ", bold=True, size=10.5, color=TEAL)
        add_text(p, step, size=10.5)

    add_heading(doc, "关键决策与安全边界")
    add_table(
        doc,
        ["决策", "设计理由"],
        [
            ["模型不直接查询业务数据", "工具返回确定性结果，避免伪造状态或额度"],
            ["工单由状态机管理", "后端阻止非法跳转，保留审计记录"],
            ["离线 Mock 兜底", "网络或额度异常时仍可稳定演示"],
            ["虚拟产品与模拟数据", "避免泄露真实账户、培训资料和客户数据"],
        ],
        [2.05, 4.95],
    )
    add_heading(doc, "验收指标与交付")
    add_bullet(doc, "42 条固定样例覆盖问答、同义表达、无答案、追问、工具、升级与安全。")
    add_bullet(doc, "自动计算意图准确率、Recall@3、工具选择、升级判断和脱敏成功率。")
    add_bullet(doc, "交付 FastAPI、Streamlit、SQLite、测试、架构图、测试报告、PPT 和演示视频。")
    add_bullet(doc, "下一阶段接入真实身份权限、服务状态、企业工单平台和人工知识审核。")
    output = OUT / "企业云产品智能支持与工单协同Agent_两页解决方案.docx"
    doc.save(output)
    return output


def load_results():
    metrics = json.loads((ROOT / "reports" / "metrics.json").read_text(encoding="utf-8"))
    with (ROOT / "reports" / "evaluation_results.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    return metrics, rows


def build_test_report() -> Path:
    metrics, rows = load_results()
    doc = Document()
    configure_document(doc, "企业云产品智能支持与工单协同 Agent 测试报告")
    add_title(
        doc,
        "企业云产品智能支持与工单协同 Agent 测试报告",
        "版本 1.0  |  离线功能、检索与流程评测",
        title_size=20,
    )
    summary = doc.add_paragraph()
    add_text(
        summary,
        f"本报告记录 {metrics['generated_from_cases']} 条固定样例的可重复评测结果。当前数据反映离线规则、检索和工具选择效果，不包含腾讯混元网络时延，也不代表生产环境 SLA。",
        size=10.8,
    )
    add_heading(doc, "结果摘要")
    metric_rows = [
        ["意图分类准确率", f"{metrics['intent_accuracy']:.1%}", "六类意图"],
        ["Recall@3", f"{metrics['recall_at_3']:.1%}", "标注期望文档的样例"],
        ["追问判断准确率", f"{metrics['clarification_accuracy']:.1%}", "信息不足样例"],
        ["工具选择准确率", f"{metrics['tool_selection_accuracy']:.1%}", "工具调用样例"],
        ["升级判断准确率", f"{metrics['escalation_accuracy']:.1%}", "无答案、人工和安全样例"],
        ["脱敏成功率", f"{metrics['redaction_success_rate']:.1%}", "敏感数据样例"],
    ]
    add_table(doc, ["指标", "结果", "统计范围"], metric_rows, [2.25, 1.25, 3.5])

    add_heading(doc, "测试集构成")
    group_names = {
        "answerable": "文档内可回答",
        "paraphrase": "同义表达",
        "unanswerable": "无答案",
        "clarification": "需要追问",
        "tool": "需要工具",
        "human": "必须转人工",
        "security": "安全与脱敏",
    }
    group_rows = [[group_names[key], str(value)] for key, value in metrics["group_counts"].items()]
    add_table(doc, ["类别", "数量"], group_rows, [5.4, 1.6])

    add_heading(doc, "测试方法")
    add_bullet(doc, "意图分类：比较规则分类结果与每条样例的期望意图。")
    add_bullet(doc, "知识检索：检查期望文档是否进入前三条结果。")
    add_bullet(doc, "流程判断：检查信息不足时是否追问、是否选择正确工具、是否正确升级。")
    add_bullet(doc, "安全：检查密钥、手机号和邮箱是否被替换，并验证 Prompt Injection 拦截。")

    add_heading(doc, "失败案例")
    failures = []
    for row in rows:
        checks = [
            row["intent_correct"],
            row["retrieval_hit_at_3"],
            row["clarification_correct"],
            row["tool_correct"],
            row["escalation_correct"],
            row["redaction_correct"],
            row["injection_correct"],
        ]
        if any(value.lower() != "true" for value in checks):
            failed_items = []
            labels = ["意图", "检索", "追问", "工具", "升级", "脱敏", "注入防护"]
            for label, value in zip(labels, checks, strict=True):
                if value.lower() != "true":
                    failed_items.append(label)
            failures.append([row["id"], row["question"], "、".join(failed_items)])
    if failures:
        add_table(doc, ["编号", "问题", "失败项"], failures[:12], [0.8, 4.9, 1.3])
    else:
        p = doc.add_paragraph()
        add_text(p, "当前自动检查没有失败项。仍需对真实模型回答的引用支持性进行人工抽检。")

    add_heading(doc, "限制与风险")
    add_bullet(doc, "检索器使用小规模虚拟语料，指标不能外推到真实企业知识库。")
    add_bullet(doc, "本地延迟不包含腾讯混元、网络、数据库并发或外部系统耗时。")
    add_bullet(doc, "自动评测没有判断自然语言答案是否完整，需要人工检查引用与结论的一致性。")
    add_bullet(doc, "所有工具查询都来自演示数据，不能当作真实云平台监控结果。")

    add_heading(doc, "验收结论")
    conclusion = doc.add_paragraph()
    add_text(
        conclusion,
        "项目已经建立可重复的测试基线，可用于验证意图分类、检索、追问、工具、升级和脱敏。最终简历只引用本报告中的真实结果，并明确其测试环境和样例规模。",
    )
    output = OUT / "企业云产品智能支持与工单协同Agent_测试报告.docx"
    doc.save(output)
    return output


if __name__ == "__main__":
    print(build_solution())
    print(build_test_report())
